import openpyxl

def getRowCount(filename,sheetname):
    workbook = openpyxl.load_workbook(filename)
    sheet = workbook[sheetname]
    return sheet.max_row

def getColumnCount(filename,sheetname):
    workbook = openpyxl.load_workbook(filename)
    sheet = workbook[sheetname]
    return sheet.max_column

def getCellData(filename, sheetname, rowNumber, columnNumber):
    workbook = openpyxl.load_workbook(filename)
    sheet = workbook[sheetname]
    return sheet.cell(row=rowNumber, column=columnNumber).value

def setCellData(filename, sheetname, rowNumber, columnNumber,inputData):
    workbook = openpyxl.load_workbook(filename)
    sheet = workbook[sheetname]
    sheet.cell(row=rowNumber, column=columnNumber).value = inputData
    workbook.save(filename)

print(getRowCount("test_data.xlsx","Sheet2"))


