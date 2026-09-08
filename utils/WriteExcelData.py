from pathlib import Path
import openpyxl

# excel_file = Path(__file__).with_name("test_data.xlsx")
workbook = openpyxl.load_workbook("test_data.xlsx")
sheet = workbook["Sheet2"]
# sheet.cell(row=1, column=3).value = "write data"
# workbook.save("test_data.xlsx")
rows = sheet.max_row
cols = sheet.max_column

for rows in range(1, rows+1):
    for cols in range(1, cols+1):
        sheet.cell(row=1, column=3).value = "write data"
workbook.save("test_data.xlsx")

